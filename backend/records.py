import io
import uuid
from datetime import datetime, timedelta, timezone
from typing import List, Optional, Literal

from fastapi import APIRouter, HTTPException, Depends
from pydantic import BaseModel, Field

from auth import get_current_user
from db import db
from security import engineer_project

router = APIRouter()
WORD_CONDS = ("high", "low", "equal", "range")


class RecordIn(BaseModel):
    number: int = Field(ge=1, le=100)
    name: str = ""
    interval_s: int = Field(5, ge=1, le=3600)
    channels: List[str] = Field(default_factory=list, max_length=99)
    enabled: bool = True


class AlarmDefIn(BaseModel):
    kind: Literal["bit", "word"]
    tag_id: str
    group: int = 1
    condition: str = "on"
    value: Optional[float] = None
    low: Optional[float] = None
    high: Optional[float] = None
    data_format: str = "INT16"
    content: str = ""
    library_id: str = ""
    record: bool = True
    not_save_off: bool = False
    beep: bool = False
    beep_once: bool = False
    alarm_screen: str = ""
    popup_once: bool = True


def retention(settings: Optional[dict]):
    s = settings or {}
    return bool(s.get("record_retention_enabled", True)), min(3650, max(1, int(s.get("record_retention_days") or 90)))


def iso(ts):
    return ts.replace(tzinfo=timezone.utc).isoformat() if ts else None


async def cleanup_project(p: dict) -> int:
    on, days = retention(p.get("settings"))
    if not on:
        return 0
    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    res = await db.record_samples.delete_many({"project_id": p["id"], "ts": {"$lt": cutoff}})
    info = {"at": datetime.now(timezone.utc).isoformat(), "deleted": res.deleted_count, "cutoff": cutoff.isoformat()}
    await db.projects.update_one({"id": p["id"]}, {"$set": {"record_cleanup": info}})
    return res.deleted_count


async def cleanup_all():
    async for p in db.projects.find({}, {"_id": 0, "id": 1, "settings": 1}):
        await cleanup_project(p)


@router.get("/projects/{pid}/records/storage")
async def record_storage(pid: str, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    p = await db.projects.find_one({"id": pid}, {"_id": 0, "settings": 1, "record_cleanup": 1})
    q = {"project_id": pid}
    n = await db.record_samples.count_documents(q)
    first = await db.record_samples.find_one(q, {"_id": 0, "ts": 1}, sort=[("ts", 1)])
    last = await db.record_samples.find_one(q, {"_id": 0, "ts": 1}, sort=[("ts", -1)])
    try:
        avg = (await db.command("collStats", "record_samples")).get("avgObjSize", 0)
    except Exception:
        avg = 0
    on, days = retention(p.get("settings"))
    return {"samples": n, "oldest": iso(first and first["ts"]), "newest": iso(last and last["ts"]), "est_bytes": int(n * avg),
            "enabled": on, "days": days, "last_cleanup": p.get("record_cleanup")}


@router.post("/projects/{pid}/records/cleanup")
async def record_cleanup_now(pid: str, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    p = await db.projects.find_one({"id": pid}, {"_id": 0, "id": 1, "settings": 1})
    on, days = retention(p.get("settings"))
    if not on:
        raise HTTPException(400, "Hapus otomatis sedang nonaktif")
    return {"deleted": await cleanup_project(p), "days": days}


async def check_tags(pid: str, ids: List[str]):
    n = await db.tags.count_documents({"project_id": pid, "id": {"$in": ids}})
    if n != len(set(ids)):
        raise HTTPException(400, "Ada channel/tag yang tidak valid")


@router.get("/projects/{pid}/records")
async def list_records(pid: str, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    return await db.data_records.find({"project_id": pid}, {"_id": 0}).sort("number", 1).to_list(100)


@router.post("/projects/{pid}/records")
async def add_record(pid: str, body: RecordIn, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    if await db.data_records.count_documents({"project_id": pid}) >= 100:
        raise HTTPException(400, "Maksimal 100 data record")
    if await db.data_records.find_one({"project_id": pid, "number": body.number}):
        raise HTTPException(400, f"Data record nomor {body.number} sudah ada")
    await check_tags(pid, body.channels)
    r = body.model_dump() | {"id": str(uuid.uuid4()), "project_id": pid}
    await db.data_records.insert_one(dict(r))
    return r


@router.put("/projects/{pid}/records/{rid}")
async def edit_record(pid: str, rid: str, body: RecordIn, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    if await db.data_records.find_one({"project_id": pid, "number": body.number, "id": {"$ne": rid}}):
        raise HTTPException(400, f"Data record nomor {body.number} sudah ada")
    await check_tags(pid, body.channels)
    await db.data_records.update_one({"id": rid, "project_id": pid}, {"$set": body.model_dump()})
    return await db.data_records.find_one({"id": rid}, {"_id": 0})


@router.delete("/projects/{pid}/records/{rid}")
async def del_record(pid: str, rid: str, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    await db.data_records.delete_one({"id": rid, "project_id": pid})
    return {"ok": True}


def validate_def(body: AlarmDefIn):
    if body.kind == "bit" and body.condition not in ("on", "off"):
        raise HTTPException(400, "Kondisi bit alarm harus ON atau OFF")
    if body.kind == "word":
        if body.condition not in WORD_CONDS:
            raise HTTPException(400, "Kondisi word alarm tidak valid")
        if body.condition == "range" and (body.low is None or body.high is None):
            raise HTTPException(400, "Range alarm butuh batas bawah & atas")
        if body.condition != "range" and body.value is None:
            raise HTTPException(400, "Nilai batas alarm wajib diisi")


@router.get("/projects/{pid}/alarm-defs")
async def list_defs(pid: str, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    return await db.alarm_defs.find({"project_id": pid}, {"_id": 0}).to_list(2000)


@router.post("/projects/{pid}/alarm-defs")
async def add_def(pid: str, body: AlarmDefIn, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    validate_def(body)
    await check_tags(pid, [body.tag_id])
    d = body.model_dump() | {"id": str(uuid.uuid4()), "project_id": pid}
    await db.alarm_defs.insert_one(dict(d))
    return d


@router.put("/projects/{pid}/alarm-defs/{did}")
async def edit_def(pid: str, did: str, body: AlarmDefIn, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    validate_def(body)
    await check_tags(pid, [body.tag_id])
    await db.alarm_defs.update_one({"id": did, "project_id": pid}, {"$set": body.model_dump()})
    return await db.alarm_defs.find_one({"id": did}, {"_id": 0})


@router.delete("/projects/{pid}/alarm-defs/{did}")
async def del_def(pid: str, did: str, user=Depends(get_current_user)):
    await engineer_project(pid, user)
    await db.alarm_defs.delete_one({"id": did, "project_id": pid})
    return {"ok": True}


def _local(ts: datetime, tz: int) -> datetime:
    return ts.replace(tzinfo=None) - timedelta(minutes=tz)


def _log_table(record: dict, tags: dict, rows: list, tz: int, max_cols: int = 99):
    chans = [c for c in record["channels"] if c in tags][:max_cols]
    head = ["Waktu"] + [f"{tags[c]['name']} ({tags[c].get('unit') or '-'})" for c in chans]
    fmt = "%d/%m/%Y %H:%M:%S"
    body = [[_local(r["ts"], tz).strftime(fmt)] + [None if r["v"].get(c) is None else round(r["v"][c], int(tags[c].get("decimals", 0) or 0)) for c in chans] for r in rows]
    return head, body


def build_log(title: str, record: dict, tags: dict, rows: list, start: datetime, end: datetime, tz: int, fmt: str):
    if fmt == "csv":
        import csv
        head, body = _log_table(record, tags, rows, tz)
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(head)
        w.writerows([["" if x is None else x for x in r] for r in body])
        return ("\ufeff" + buf.getvalue()).encode("utf-8"), "text/csv"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        head, body = _log_table(record, tags, rows, tz)
        wb = Workbook()
        ws = wb.active
        ws.title = f"Record {record['number']}"
        ws.append([f"{title} — Data Record #{record['number']} {record.get('name', '')}"])
        ws.append([f"{_local(start, tz):%d/%m/%Y %H:%M:%S} s/d {_local(end, tz):%d/%m/%Y %H:%M:%S} · {len(rows)} sampel"])
        ws.append(head)
        for c in ws[3]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E293B")
        for r in body:
            ws.append(r)
        ws.column_dimensions["A"].width = 21
        ws.freeze_panes = "B4"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    head, body = _log_table(record, tags, rows, tz, 12)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1 * cm, bottomMargin=1 * cm)
    st = getSampleStyleSheet()
    t = Table([head] + [["" if x is None else x for x in r] for r in body[:20000]], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")])]))
    doc.build([Paragraph(f"<b>{title}</b> — Data Log Record #{record['number']} {record.get('name', '')}", st["Title"]),
               Paragraph(f"Rentang: {_local(start, tz):%d/%m/%Y %H:%M:%S} s/d {_local(end, tz):%d/%m/%Y %H:%M:%S} · {len(rows)} sampel", st["Normal"]),
               Spacer(1, 6), t])
    return buf.getvalue(), "application/pdf"


def build_chart_pdf(title: str, record: dict, tags: dict, rows: list, start: datetime, end: datetime, tz: int = 0) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image

    chans = [c for c in record["channels"] if c in tags]
    fig, ax = plt.subplots(figsize=(11, 5.6), dpi=120)
    xs = [_local(r["ts"], tz) for r in rows]
    for c in chans:
        ax.plot(xs, [r["v"].get(c) for r in rows], label=tags[c]["name"], linewidth=1.4)
    ax.grid(True, alpha=0.3)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M:%S" if (end - start).total_seconds() <= 86400 else "%d/%m %H:%M"))
    ax.set_title(f"History Trend — Record #{record['number']} {record.get('name', '')}")
    if chans:
        ax.legend(fontsize=7, ncol=min(len(chans), 6), loc="upper left")
    fig.autofmt_xdate()
    img = io.BytesIO()
    fig.savefig(img, format="png", bbox_inches="tight")
    plt.close(fig)
    img.seek(0)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1 * cm, bottomMargin=1 * cm)
    st = getSampleStyleSheet()
    fmt = "%d/%m/%Y %H:%M:%S"
    doc.build([Paragraph(f"<b>{title}</b> — History Trend Record #{record['number']} {record.get('name', '')}", st["Title"]),
               Paragraph(f"Rentang: {_local(start, tz).strftime(fmt)} s/d {_local(end, tz).strftime(fmt)} · {len(rows)} sampel", st["Normal"]),
               Spacer(1, 6), Image(img, width=26 * cm, height=13.2 * cm)])
    return buf.getvalue()
